from __future__ import annotations

from dataclasses import dataclass
from datetime import date, datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

@dataclass
class PersonRecord:
    surname: str
    name: str
    father: str
    birth: Any
    phone: str
    address: str

    @property
    def full_name(self) -> str:
        return " ".join(x for x in (self.surname, self.name, self.father) if x).strip()

    @property
    def birth_text(self) -> str:
        value = self.birth
        if isinstance(value, (datetime, date)):
            return value.strftime("%d.%m.%Y")
        text = str(value or "").strip()
        for fmt in ("%Y-%m-%d", "%d.%m.%Y", "%d/%m/%Y"):
            try:
                return datetime.strptime(text, fmt).strftime("%d.%m.%Y")
            except ValueError:
                pass
        return text

class ExcelRepository:
    def __init__(self) -> None:
        self.records: list[PersonRecord] = []
        self.analysis_records: list[str] = []
        self.people_file: Path | None = None
        self.analysis_file: Path | None = None

    def load_people(self, filename: str) -> tuple[int, list[str]]:
        path = Path(filename)
        wb = load_workbook(path, data_only=True, read_only=True)
        ws = wb.active
        loaded: list[PersonRecord] = []
        warnings: list[str] = []
        for index, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
            values = list(row[:6]) + [None] * max(0, 6 - len(row))
            if not any(v not in (None, "") for v in values):
                continue
            record = PersonRecord(
                surname=str(values[0] or "").strip(),
                name=str(values[1] or "").strip(),
                father=str(values[2] or "").strip(),
                birth=values[3],
                phone=str(values[4] or "").strip(),
                address=str(values[5] or "").strip(),
            )
            missing = [name for name, val in (("прізвище", record.surname), ("ім’я", record.name), ("телефон", record.phone)) if not val]
            if missing:
                warnings.append(f"Рядок {index}: немає {', '.join(missing)}")
            loaded.append(record)
        self.records = loaded
        self.people_file = path
        return len(loaded), warnings

    def load_analysis(self, filename: str) -> int:
        path = Path(filename)
        wb = load_workbook(path, data_only=True, read_only=True)
        ws = wb.active
        phones: list[str] = []
        for row in ws.iter_rows(values_only=True):
            if not row or row[0] is None:
                continue
            value = str(row[0]).strip().replace("+", "").replace(" ", "").replace("-", "")
            if "E+" in value.upper():
                try:
                    value = str(int(float(value)))
                except ValueError:
                    continue
            if len(value) >= 10:
                phones.append(value)
        self.analysis_records = phones
        self.analysis_file = path
        return len(phones)

repo = ExcelRepository()
