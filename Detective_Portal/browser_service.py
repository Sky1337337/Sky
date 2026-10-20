from __future__ import annotations

import datetime as dt
import json
import queue
import shutil
import subprocess
import threading
import time
from http.client import HTTPConnection, HTTPException
from pathlib import Path
from typing import Callable

from openpyxl import Workbook, load_workbook
from selenium import webdriver
from selenium.common.exceptions import WebDriverException, TimeoutException
from selenium.webdriver.support.ui import WebDriverWait
from selenium.webdriver.chrome.options import Options
from selenium.webdriver.chrome.service import Service
from selenium.webdriver.common.by import By
from selenium.webdriver.common.keys import Keys
from webdriver_manager.chrome import ChromeDriverManager

from app_state import SCREENSHOT_DIR, settings
from excel_data import PersonRecord, repo

LogFn = Callable[[str, str], None]

class BrowserService:
    def __init__(self) -> None:
        self.driver = None
        self.current = 0
        self.analysis_current = 0
        self._chrome_process: subprocess.Popen | None = None
        self._chrome_port: int | None = None

    PORTAL_HOST = "detective.in.np.gov.ua"
    CONNECT_TIMEOUT = 30.0
    CONNECT_POLL_INTERVAL = 0.5
    DEBUG_PROBE_TIMEOUT = 1.0

    def _select_portal_tab(self) -> bool:
        """Recover after a Chrome tab closes; never write into an unrelated tab."""
        if self.driver is None:
            return False
        try:
            handles = self.driver.window_handles
        except WebDriverException:
            return False
        for handle in handles:
            try:
                self.driver.switch_to.window(handle)
                if self.PORTAL_HOST in (self.driver.current_url or "").lower():
                    return True
            except WebDriverException:
                continue
        return False

    @property
    def connected(self) -> bool:
        return self._select_portal_tab()

    def _require_portal(self) -> None:
        if self._select_portal_tab():
            return
        # An existing Selenium session can become invalid after its Chrome tab
        # closes. Reattach once to the already running debugging Chrome.
        if self.driver is not None:
            try:
                options = Options()
                options.debugger_address = f"127.0.0.1:{settings.chrome_debug_port}"
                self.driver = webdriver.Chrome(
                    service=Service(ChromeDriverManager().install()), options=options
                )
            except WebDriverException:
                pass
        if not self._select_portal_tab():
            raise RuntimeError(
                "Не знайдено відкритої вкладки detective.in.np.gov.ua "
                "у Chrome, підключеному до програми. Відкрий портал у цьому "
                "Chrome і повтори; поточний рядок Excel не пропущено."
            )

    @staticmethod
    def find_chrome() -> str | None:
        candidates = [
            shutil.which("chrome"),
            shutil.which("chrome.exe"),
            r"C:\Program Files\Google\Chrome\Application\chrome.exe",
            r"C:\Program Files (x86)\Google\Chrome\Application\chrome.exe",
            str(Path.home() / "AppData/Local/Google/Chrome/Application/chrome.exe"),
        ]
        return next((p for p in candidates if p and Path(p).exists()), None)

    @staticmethod
    def _debugger_ready(port: int, timeout: float) -> bool:
        """Probe DevTools directly on loopback, without system HTTP proxies."""
        connection = HTTPConnection("127.0.0.1", port, timeout=timeout)
        try:
            connection.request("GET", "/json/version")
            response = connection.getresponse()
            if response.status != 200:
                return False
            # A listening TCP port (or an unrelated HTTP server) is not enough.
            payload = json.loads(response.read(65536))
            if not isinstance(payload, dict):
                return False
            browser_name = payload.get("Browser", "")
            websocket = payload.get("webSocketDebuggerUrl", "")
            return (
                isinstance(browser_name, str)
                and browser_name.startswith(("Chrome/", "Chromium/"))
                and isinstance(websocket, str)
                and websocket.startswith("ws://")
            )
        except (OSError, HTTPException, ValueError):
            return False
        finally:
            connection.close()

    @staticmethod
    def _connection_step(action: Callable, timeout: float, cleanup: Callable | None = None):
        """Bound a blocking dependency call; discard and clean up late results."""
        if timeout <= 0:
            raise TimeoutError("Час очікування підключення вичерпано")
        outcome: queue.Queue = queue.Queue(maxsize=1)
        abandoned = threading.Event()

        def clean_up():
            try:
                if cleanup is not None:
                    cleanup()
            except Exception:
                pass

        def run():
            try:
                result = action()
            except Exception as exc:
                clean_up()
                outcome.put((None, exc))
            else:
                outcome.put((result, None))
            finally:
                if abandoned.is_set():
                    clean_up()

        deadline = time.monotonic() + timeout
        threading.Thread(target=run, daemon=True).start()
        try:
            result, error = outcome.get(timeout=max(0, deadline - time.monotonic()))
            if time.monotonic() >= deadline:
                raise queue.Empty
        except queue.Empty:
            abandoned.set()
            # Cleanup must not extend the deadline either. The worker also
            # cleans up if ChromeDriver only finishes starting after timeout.
            threading.Thread(target=clean_up, daemon=True).start()
            raise TimeoutError("Час очікування підключення вичерпано") from None
        if error is not None:
            raise error
        return result

    def _probe_debugger(self, port: int, timeout: float) -> bool:
        try:
            return self._connection_step(
                lambda: self._debugger_ready(port, timeout), timeout
            )
        except TimeoutError:
            return False

    def connect(self, log: LogFn) -> None:
        port = settings.chrome_debug_port
        address = f"127.0.0.1:{port}"
        # Drop a stale attachment without closing the user's debugging Chrome.
        old_driver, self.driver = self.driver, None
        if old_driver is not None:
            try:
                self._connection_step(old_driver.service.stop, self.DEBUG_PROBE_TIMEOUT)
            except (OSError, WebDriverException):
                pass

        log(f"Готую ChromeDriver для підключення (до {self.CONNECT_TIMEOUT:g} с)…", "info")
        try:
            driver_path = self._connection_step(
                lambda: ChromeDriverManager().install(), self.CONNECT_TIMEOUT
            )
        except Exception as exc:
            raise RuntimeError(
                f"Не вдалося підготувати ChromeDriver за {self.CONNECT_TIMEOUT:g} с. Перевір доступ до мережі "
                "та повтори підключення."
            ) from exc

        deadline = time.monotonic() + self.CONNECT_TIMEOUT
        ready = self._probe_debugger(port, min(self.DEBUG_PROBE_TIMEOUT, self.CONNECT_TIMEOUT))
        if ready:
            log(f"Використовую вже запущений Chrome ({address})…", "info")
        elif (
            self._chrome_process is not None
            and self._chrome_port == port
            and self._chrome_process.poll() is None
        ):
            log("Chrome ще запускається; очікую без запуску іншого екземпляра…", "info")
        elif time.monotonic() < deadline:
            chrome = self.find_chrome()
            if not chrome:
                raise RuntimeError("Google Chrome не знайдено. Встанови Chrome та повтори підключення.")
            profile = Path.home() / "AppData/Local/DetectivePortalPro/ChromeProfile"
            profile.mkdir(parents=True, exist_ok=True)
            try:
                self._chrome_process = subprocess.Popen([
                    chrome,
                    f"--remote-debugging-port={port}",
                    f"--user-data-dir={profile}",
                    "--no-first-run",
                ])
            except OSError as exc:
                raise RuntimeError(f"Не вдалося запустити Google Chrome: {exc}") from exc
            self._chrome_port = port
            log("Запускаю Chrome у режимі віддаленого керування…", "info")

        log(f"Очікую підключення Chrome ({address}), до {self.CONNECT_TIMEOUT:g} с…", "info")
        options = Options()
        options.debugger_address = address
        last_error = None
        attempts = 0
        while time.monotonic() < deadline:
            if not ready:
                ready = self._probe_debugger(
                    port, min(self.DEBUG_PROBE_TIMEOUT, deadline - time.monotonic())
                )
            if time.monotonic() >= deadline:
                break
            if ready:
                attempts += 1
                log(f"Chrome готовий. Підключаю Selenium (спроба {attempts})…", "info")
                # A failed constructor may leave a service behind. Each retry
                # gets a fresh ChromeDriver, but never another Chrome process.
                service = Service(driver_path)
                try:
                    self.driver = self._connection_step(
                        lambda: webdriver.Chrome(service=service, options=options),
                        deadline - time.monotonic(),
                        # Stop ChromeDriver, keeping debugging Chrome and tabs.
                        cleanup=service.stop,
                    )
                except TimeoutError as exc:
                    last_error = exc
                    break
                except WebDriverException as exc:
                    last_error = exc
                    log("Selenium поки не підключився; повторюю в межах часу очікування…", "warning")
                else:
                    log("Chrome підключено. Перед внесенням відкрий detective.in.np.gov.ua у цьому Chrome.", "success")
                    return
            remaining = deadline - time.monotonic()
            if remaining <= 0:
                break
            time.sleep(min(self.CONNECT_POLL_INTERVAL, remaining))
            ready = False

        detail = "Selenium не зміг приєднатися" if last_error else "відладочний порт Chrome не готовий"
        last_detail = ""
        if last_error is not None:
            reason = getattr(last_error, "msg", None) or str(last_error)
            reason = " ".join(reason.split())[:240]
            last_detail = f" Остання помилка Selenium: {reason}."
        raise RuntimeError(
            f"Не вдалося підключити Chrome ({address}): {detail} "
            f"за {self.CONNECT_TIMEOUT:g} с; спроб Selenium: {attempts}.{last_detail} "
            "Перевір Chrome і порт у налаштуваннях та натисни «Підключити Chrome» ще раз."
        ) from last_error

    def _inputs(self):
        return self.driver.find_elements(By.TAG_NAME, "input") if self.connected else []

    def _find_field(self, words: list[str]):
        for field in self._inputs():
            try:
                haystack = " ".join(field.get_attribute(x) or "" for x in ("placeholder", "aria-label", "name", "id")).lower()
                if any(word.lower() in haystack for word in words):
                    return field
            except WebDriverException:
                continue
        return None

    @staticmethod
    def _write(field, value: str) -> bool:
        if field is None:
            return False
        field.click()
        field.clear()
        field.send_keys(str(value))
        return True

    def _arm_save_confirmation(self) -> None:
        """Observe a *new* success event, even when the old toast stays visible.

        An unchanged old green message must never count as a fresh save.
        """
        self.driver.execute_script(r"""
            if (window.__dpSaveObserver) window.__dpSaveObserver.disconnect();
            window.__dpNewSaveAccepted = false;
            const success = (s) => /Особу\s+(прийнято|внесено)/i.test(s || '');
            const observer = new MutationObserver((mutations) => {
                for (const m of mutations) {
                    if (m.type === 'characterData') {
                        const p = m.target.parentElement;
                        if (p && success(p.textContent)) window.__dpNewSaveAccepted = true;
                    }
                    if (m.type === 'childList') {
                        for (const node of m.addedNodes) {
                            if (success(node.textContent)) window.__dpNewSaveAccepted = true;
                        }
                    }
                }
            });
            observer.observe(document.body, {childList:true, subtree:true, characterData:true});
            window.__dpSaveObserver = observer;
        """)

    def _wait_accepted(self, timeout: int = 45) -> None:
        """Only a success DOM update *after saving* confirms this row."""
        try:
            WebDriverWait(self.driver, timeout, poll_frequency=0.2).until(
                lambda driver: driver.execute_script(
                    "return window.__dpNewSaveAccepted === true;"
                )
            )
        except TimeoutException as exc:
            raise RuntimeError(
                f"За {timeout} с не отримано НОВОГО підтвердження «Особу прийнято» "
                "після збереження. Стара зелена плашка не зараховується. "
                "Перевір запис на порталі перед повторним запуском."
            ) from exc
        finally:
            try:
                self.driver.execute_script(
                    "if (window.__dpSaveObserver) window.__dpSaveObserver.disconnect();"
                )
            except WebDriverException:
                pass

    def _click_new_person(self, log: LogFn) -> None:
        """Select the explicit New person control after successful save."""
        xpath = ("//*[self::button or self::a or @role='button' or contains(@class,'dx-button')]"
                 "[.//*[normalize-space()='Нова особа'] or normalize-space()='Нова особа']")
        def find_button(driver):
            for button in driver.find_elements(By.XPATH, xpath):
                try:
                    if button.is_displayed() and button.is_enabled():
                        return button
                except WebDriverException:
                    continue
            return False
        button = WebDriverWait(self.driver, 12, poll_frequency=0.25).until(find_button)
        button.click()
        log('Після підтвердження натиснуто «Нова особа»', 'success')

    def fill_record(self, record: PersonRecord, log: LogFn) -> int:
        self._require_portal()
        mappings = [
            (["прізвище"], record.surname),
            (["ім'я", "імя"], record.name),
            (["по батькові", "батькові"], record.father),
            (["дата", "народ"], record.birth_text),
            (["адреса"], record.address),
            (["від 9 до 13 цифр", "телефон", "phone"], record.phone),
        ]
        found = 0

        # Special DevExpress date field handling
        try:
            if record.birth_text:
                date_field = self.driver.find_element(
                    By.CSS_SELECTOR,
                    "input.dx-texteditor-input[placeholder='дд.мм.рррр']"
                )
                date_field.click()
                date_field.send_keys(Keys.CONTROL + "a")
                date_field.send_keys(Keys.BACKSPACE)
                date_field.send_keys(record.birth_text)
                date_field.send_keys(Keys.TAB)
                found += 1
                time.sleep(settings.action_delay)
        except Exception:
            pass

        for aliases, value in mappings:
            if self._write(self._find_field(aliases), value):
                found += 1
                time.sleep(settings.action_delay)
        # Saving is not considered successful until the server confirms it.
        # Avoid silently swallowing missing save controls and advancing Excel.
        save_xpath = ("//*[self::button or @role='button' or contains(@class,'dx-button')]"
                      "[.//*[normalize-space()='Зберегти'] or normalize-space()='Зберегти']")
        button = WebDriverWait(self.driver, 10).until(
            lambda d: next((el for el in d.find_elements(By.XPATH, save_xpath)
                            if el.is_displayed() and el.is_enabled()), False)
        )
        # Start watching BEFORE clicking save; a persistent previous toast is OK.
        self._arm_save_confirmation()
        button.click()
        log('Натиснуто «Зберегти»; очікую підтвердження...', 'info')
        self._wait_accepted()
        log('Отримано підтвердження «Особу прийнято»', 'success')
        self._click_new_person(log)
        return found

    def screenshot(self, prefix: str = "error") -> Path | None:
        if not self._select_portal_tab():
            return None
        path = SCREENSHOT_DIR / f"{prefix}_{dt.datetime.now():%Y%m%d_%H%M%S}.png"
        try:
            self.driver.save_screenshot(str(path))
            return path
        except Exception:
            return None

    def analyze_phone(self, phone: str) -> tuple[list[str], Path]:
        self._require_portal()
        sim_input = self.driver.find_element(By.CSS_SELECTOR, 'input[placeholder="Введіть SIM"]')
        sim_input.click(); sim_input.clear(); sim_input.send_keys(phone)
        self.driver.find_element(By.CSS_SELECTOR, 'div[aria-label="Пошук"]').click()
        time.sleep(2)
        rows = [row.text.strip() for row in self.driver.find_elements(By.CSS_SELECTOR, ".dx-datagrid-rowsview .dx-data-row") if row.text.strip()]
        if not rows:
            rows = [self.driver.find_element(By.TAG_NAME, "body").text]
        out_dir = Path(settings.result_folder)
        out_dir.mkdir(parents=True, exist_ok=True)
        output = out_dir / "analysis_all.xlsx"
        if output.exists():
            wb = load_workbook(output)
            ws = wb.active
        else:
            wb = Workbook(); ws = wb.active; ws.title = "Analysis"; ws.append(["Телефон", "Дата", "Результат"])
        ws.append([phone, dt.datetime.now().strftime("%d.%m.%Y %H:%M:%S"), "\n".join(rows)])
        ws.column_dimensions["A"].width = 18; ws.column_dimensions["B"].width = 22; ws.column_dimensions["C"].width = 80
        wb.save(output)
        return rows, output

browser = BrowserService()
